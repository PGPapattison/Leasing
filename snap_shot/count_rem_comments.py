#!/usr/bin/env python3
"""Open each recovered xlsx in recovery_out/ and count non-empty REM ClickUp
Comment cells (column P in property blocks). Print a summary."""
import os, sys, glob
from openpyxl import load_workbook

def count_rem_comments(path):
    wb = load_workbook(path, data_only=True, read_only=False)
    total = 0
    per_sheet = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        n = 0
        # Sweep column P (16), skip the header text 'REM ClickUp Comment'
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row=row, column=16).value
            if v and isinstance(v, str) and v.strip() and v.strip().upper() != "REM CLICKUP COMMENT":
                n += 1
        if n:
            per_sheet[sn] = n
            total += n
    wb.close()
    return total, per_sheet

def main():
    files = sorted(glob.glob("recovery_out/snap_v*.xlsx"))
    print(f"Found {len(files)} recovered files\n")
    results = []
    for f in files:
        try:
            total, per = count_rem_comments(f)
            results.append((f, total, per))
            print(f"  {os.path.basename(f)}")
            print(f"    total REM Comments: {total}")
            if per:
                for sn, n in per.items():
                    print(f"      {sn}: {n}")
            print()
        except Exception as e:
            print(f"  {os.path.basename(f)}: ERROR {e}\n")
    print("=" * 60)
    print("BEST CANDIDATE (most REM Comments):")
    if results:
        best = max(results, key=lambda x: x[1])
        print(f"  {os.path.basename(best[0])} → {best[1]} REM Comments")

if __name__ == "__main__":
    main()
