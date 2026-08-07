"""Dump rows around each REM comment location for debugging."""
import glob
from openpyxl import load_workbook

path = sorted(glob.glob("recovery_out/snap_v02_*.xlsx"))[0]
wb = load_workbook(path, data_only=True)
ws = wb["Snap Shot"]

for target in [97, 101, 1303, 1306]:
    print(f"\n{'='*70}\nContext around row {target}:")
    for r in range(max(1, target - 15), target + 3):
        vals = [ws.cell(row=r, column=c).value for c in range(2, 8)]  # B..G
        vals = [str(v)[:30] if v else "" for v in vals]
        marker = " <-- REM COMMENT" if r == target else ""
        print(f"  row {r:4d}: {' | '.join(vals)}{marker}")
